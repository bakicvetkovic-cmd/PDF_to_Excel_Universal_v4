import tkinter as tk
from tkinter import ttk, filedialog, messagebox, simpledialog
from pathlib import Path
import json, re, sys, subprocess

# V4: no Pillow. PDF is read directly with PyMuPDF.
# Excel workbook is created with openpyxl.

try:
    import fitz
except ImportError:
    fitz=None
try:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
except ImportError:
    Workbook=None

TEMPLATES=Path.home()/".pdf_excel_universal"/"templates"
TEMPLATES.mkdir(parents=True, exist_ok=True)

def num(v):
    v=v.strip().replace(" ","")
    if not v: return None
    if re.fullmatch(r"-?\d{1,3}(?:\.\d{3})*(?:,\d+)?",v) or re.fullmatch(r"-?\d+(?:,\d+)?",v):
        try: return float(v.replace(".","").replace(",","."))
        except: pass
    return v

class App:
    def __init__(self,root):
        self.root=root; root.title("PDF to Excel Universal Converter v4")
        root.geometry("1250x820")
        self.doc=None; self.path=None; self.page=0
        self.columns=[]; self.start_y=0
        self.page_columns=None  # independent column widths per page
        self.page_start_y=None  # independent data start (Y) per page
        self.split_start_y_by_page=False
        self.split_column_widths=False
        self.zoom=1.0; self.drag=None; self.photo=None
        self.build()

    def build(self):
        bar=ttk.Frame(self.root,padding=6); bar.pack(fill="x")
        for txt,cmd in [
            ("Open PDF",self.open),
            ("◀",self.prev),("▶",self.next),
            ("Add Column",self.add),("Delete Column",self.delete),
            ("Rename",self.rename),("Save Template",self.save),
            ("Load Template",self.load),("Separate Column Widths by Page",self.split_widths),("Separate Data Start Y by Page",self.split_start_y),("EXPORT TO EXCEL",self.export)]:
            ttk.Button(bar,text=txt,command=cmd).pack(side="left",padx=2)
        self.info=ttk.Label(bar,text="No PDF open"); self.info.pack(side="right")

        pane=ttk.Panedwindow(self.root,orient="horizontal"); pane.pack(fill="both",expand=True,padx=6,pady=4)
        lf=ttk.Frame(pane); rf=ttk.Frame(pane,padding=8); pane.add(lf,weight=5); pane.add(rf,weight=2)
        self.canvas=tk.Canvas(lf,bg="#777"); self.canvas.pack(fill="both",expand=True)
        self.canvas.bind("<Button-1>",self.click); self.canvas.bind("<B1-Motion>",self.drag_move); self.canvas.bind("<ButtonRelease-1>",lambda e:setattr(self,"drag",None))
        self.tree=ttk.Treeview(rf,columns=("name","x1","x2"),show="headings")
        for c,t,w in [("name","Column",170),("x1","From",65),("x2","To",65)]:
            self.tree.heading(c,text=t); self.tree.column(c,width=w)
        self.tree.pack(fill="both",expand=True)
        ttk.Button(rf,text="Data Start (Y)",command=self.start).pack(fill="x",pady=3)
        ttk.Button(rf,text="Add Empty Column",command=self.add).pack(fill="x",pady=3)
        ttk.Label(rf,text="How to Use",font=("Segoe UI",10,"bold")).pack(anchor="w",pady=(15,4))
        ttk.Label(rf,text=
          "Red lines = column boundaries.\n"
          "Drag a red line to adjust a column boundary.\n"
          "Blue line = where data extraction begins.\n\n"
          "Templates are saved as JSON and can be reused for PDFs with the same layout.\n\n"
          "The “Separate Column Widths by Page” button lets each page use its own column widths.\n\n"
          "V4 does not use Pillow. Only PyMuPDF and openpyxl are required."
        ,wraplength=260,justify="left").pack(anchor="w")
        self.status=tk.StringVar(value="Ready.")
        ttk.Label(self.root,textvariable=self.status,relief="sunken",anchor="w").pack(fill="x",side="bottom")

    def open(self):
        if fitz is None:
            messagebox.showerror("PyMuPDF Missing","Run INSTALL.bat included in this package.")
            return
        p=filedialog.askopenfilename(filetypes=[("PDF","*.pdf")])
        if not p:return
        self.doc=fitz.open(p); self.path=p; self.page=0
        self.page_columns=None
        self.page_start_y=None
        self.split_column_widths=False
        w=self.doc[0].rect.width
        # Starting layout based on the supplied accounting-card PDF.
        cuts=[55,105,150,220,280,360,395,450,505,565,620]
        names=["Entry Date","Entry No.","Document Date","Document No.","Due Date","Description","Status","Debit","Credit","Balance","Account"]
        x=0; self.columns=[]
        for name,c in zip(names,cuts):
            self.columns.append({"name":name,"x1":x,"x2":min(c,w)}); x=min(c,w)
        self.start_y=150
        self.draw()
        self.status.set(f"Opened {Path(p).name} | {len(self.doc)} pages")
    def prev(self):
        if self.doc and self.page>0:
            self.page-=1;self.draw()
            if self.split_column_widths or getattr(self, "split_start_y_by_page", False):self.status.set(f"Page-specific setup: page {self.page+1}/{len(self.doc)}")
    def next(self):
        if self.doc and self.page+1<len(self.doc):
            self.page+=1;self.draw()
            if self.split_column_widths or getattr(self, "split_start_y_by_page", False):self.status.set(f"Page-specific setup: page {self.page+1}/{len(self.doc)}")

    def draw(self):
        if not self.doc:return
        page=self.doc[self.page]
        active_columns = self.get_active_columns()
        # Render PDF using Tk's PhotoImage via PPM data; no Pillow.
        pix=page.get_pixmap(matrix=fitz.Matrix(self.zoom,self.zoom),alpha=False)
        ppm=pix.tobytes("ppm")
        self.photo=tk.PhotoImage(data=ppm)
        self.canvas.delete("all")
        self.canvas.create_image(10,10,image=self.photo,anchor="nw")
        for i,c in enumerate(active_columns):
            x=10+c["x2"]*self.zoom
            self.canvas.create_line(x,10,x,1000,fill="red",width=2)
            self.canvas.create_text(x+3,14,text=str(i+1),fill="red",anchor="nw")
        y=10+self.get_active_start_y()*self.zoom
        self.canvas.create_line(10,y,1200,y,fill="blue",width=2)
        self.canvas.create_text(14,y+3,text="DATA START",fill="blue",anchor="nw")
        self.refresh()
        self.info.config(text=f"Page {self.page+1}/{len(self.doc)}")

    def get_active_columns(self):
        if self.split_column_widths and self.page_columns is not None:
            return self.page_columns[self.page]
        return self.columns

    def get_active_columns_for_page(self, page_index):
        if self.split_column_widths and self.page_columns is not None:
            if 0 <= page_index < len(self.page_columns):
                return self.page_columns[page_index]
        return self.columns

    def get_active_start_y(self):
        if self.split_start_y_by_page and self.page_start_y is not None:
            return self.page_start_y[self.page]
        return self.start_y

    def get_start_y_for_page(self, page_index):
        if self.split_start_y_by_page and self.page_start_y is not None:
            if 0 <= page_index < len(self.page_start_y):
                return self.page_start_y[page_index]
        return self.start_y

    def split_start_y(self):
        if not self.doc:
            return
        if getattr(self, "split_start_y_by_page", False):
            messagebox.showinfo(
                "Data Start Y Already Separated",
                "Data Start (Y) is already separated by page.\n"
                "You can now adjust the data start separately for each page."
            )
            return
        self.page_start_y = [self.start_y for _ in range(len(self.doc))]
        self.split_start_y_by_page = True
        self.draw()
        self.status.set(
            "Data start Y is now independent for each page. "
            f"Adjusting page {self.page+1}/{len(self.doc)}."
        )

    def split_widths(self):
        if not self.doc:
            return
        if self.split_column_widths:
            messagebox.showinfo(
                "Column Widths Already Separated",
                "Column widths are already separated by page.\n"
                "You can now change widths on one page without affecting the others."
            )
            return
        self.page_columns = [
            [dict(c) for c in self.columns] for _ in range(len(self.doc))
        ]
        self.split_column_widths = True
        self.refresh()
        self.draw()
        self.status.set(
            "Column widths are now independent for each page. "
            f"Adjusting page {self.page+1}/{len(self.doc)}."
        )

    def release(self,e):
        self.drag=None

    def refresh(self):
        for i in self.tree.get_children():self.tree.delete(i)
        for i,c in enumerate(self.get_active_columns()):
            self.tree.insert("", "end",iid=str(i),values=(c["name"],round(c["x1"],1),round(c["x2"],1)))

    def click(self,e):
        # The blue line is always available to move.
        # Before separation, this changes Y for all pages; afterward
        # it changes only the active page.
        active_y=self.get_active_start_y()
        y_line=10+active_y*self.zoom
        if abs(e.y-y_line) <= 15:
            self.drag=("y",)
            return

        active_columns=self.get_active_columns()
        for i,c in enumerate(active_columns):
            x=10+c["x2"]*self.zoom
            if abs(e.x-x)<=8:
                self.drag=("col",i)
                return
        self.drag=None

    def drag_move(self,e):
        if not self.drag:
            return

        # Y movement: before separation, shared across pages;
        # after separation, independent for the active page.
        if self.drag[0] == "y":
            y=max(0,(e.y-10)/self.zoom)
            y=min(y,self.doc[self.page].rect.height-1)
            if self.split_start_y_by_page and self.page_start_y is not None:
                self.page_start_y[self.page]=y
            else:
                self.start_y=y
            self.draw()
            return

        # Existing column movement logic — behavior unchanged.
        if self.drag[0] == "col":
            i=self.drag[1]
            cols=self.get_active_columns()
            x=(e.x-10)/self.zoom
            if i>0:
                x=max(x,cols[i-1]["x2"]+1)
            x=min(x,self.doc[self.page].rect.width-1)
            cols[i]["x2"]=x
            if i+1<len(cols):
                cols[i+1]["x1"]=x
            self.draw()

    def add(self):
        if not self.doc:return
        name=simpledialog.askstring("Column","Name of the new column:")
        if not name:return
        width=self.doc[self.page].rect.width
        cols=self.get_active_columns()
        x1=cols[-1]["x2"] if cols else 0
        x2=min(width,x1+60)
        newcol={"name":name,"x1":x1,"x2":x2}
        cols.append(newcol)
        if self.split_column_widths:
            for pgcols in self.page_columns:
                if len(pgcols) < len(cols):
                    px1=pgcols[-1]["x2"] if pgcols else 0
                    pgcols.append({"name":name,"x1":px1,
                                   "x2":min(self.doc[0].rect.width,px1+60)})
                else:
                    pgcols[-1]["name"]=name
        self.draw()

    def delete(self):
        s=self.tree.selection()
        if not s:return
        i=int(s[0])
        cols=self.get_active_columns()
        cols.pop(i)
        if i<len(cols):cols[i]["x1"]=cols[i-1]["x2"] if i else 0
        if self.split_column_widths:
            for pgcols in self.page_columns:
                if i < len(pgcols):
                    pgcols.pop(i)
                if i < len(pgcols):
                    pgcols[i]["x1"]=pgcols[i-1]["x2"] if i else 0
        else:
            self.columns=cols
        self.draw()

    def rename(self):
        s=self.tree.selection()
        if not s:return
        i=int(s[0])
        cols=self.get_active_columns()
        n=simpledialog.askstring("Column","New name:",initialvalue=cols[i]["name"])
        if n:
            if self.split_column_widths:
                for pgcols in self.page_columns:
                    pgcols[i]["name"]=n
            else:
                self.columns[i]["name"]=n
            self.draw()

    def start(self):
        if not self.doc:return
        y=simpledialog.askfloat("Data Start","Y coordinate in the PDF:",initialvalue=self.start_y,minvalue=0)
        if y is not None:self.start_y=y;self.draw()

    def save(self):
        if not self.columns:return
        n=simpledialog.askstring("Template","Template name:")
        if not n:return
        p=TEMPLATES/(re.sub(r"[^A-Za-z0-9_-]+","_",n)+".json")
        data={"name":n,"columns":self.columns,"start_y":self.start_y}
        if self.split_column_widths and self.page_columns is not None:
            data["split_column_widths"]=True
            data["page_columns"]=self.page_columns
        if getattr(self, "split_start_y_by_page", False) and self.page_start_y is not None:
            data["split_start_y_by_page"]=True
            data["page_start_y"]=self.page_start_y
        p.write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding="utf8")
        self.status.set("Saved template: "+p.name)

    def load(self):
        p=filedialog.askopenfilename(initialdir=str(TEMPLATES),filetypes=[("JSON template","*.json")])
        if not p:return
        d=json.loads(Path(p).read_text(encoding="utf8"))
        self.columns=d["columns"];self.start_y=d.get("start_y",0)
        if d.get("split_column_widths") and self.doc:
            saved=d.get("page_columns") or []
            self.page_columns=[[dict(c) for c in cols] for cols in saved[:len(self.doc)]]
            while len(self.page_columns)<len(self.doc):
                self.page_columns.append([dict(c) for c in self.columns])
            self.split_column_widths=True
        else:
            self.page_columns=None
            self.split_column_widths=False
        if d.get("split_start_y_by_page") and self.doc:
            saved_y=d.get("page_start_y") or []
            self.page_start_y=list(saved_y[:len(self.doc)])
            while len(self.page_start_y)<len(self.doc):
                self.page_start_y.append(self.start_y)
            self.split_start_y_by_page=True
        else:
            self.page_start_y=None
            self.split_start_y_by_page=False
        self.draw()
        self.status.set("Loaded: "+Path(p).name)

    def rows(self,page, page_index=None):
        if page_index is None:
            page_index = self.page
        active_columns=self.get_active_columns_for_page(page_index)
        words=page.get_text("words")
        rows=[]
        for w in sorted(words,key=lambda z:(z[1],z[0])):
            if w[1]<self.get_start_y_for_page(page_index):continue
            cy=(w[1]+w[3])/2
            target=None
            for r in rows:
                if abs(r[0]-cy)<=3.5:target=r;break
            if target:target[1].append(w)
            else:rows.append([cy,[w]])
        out=[]
        for _,ws in rows:
            ws.sort(key=lambda z:z[0]); cells=[]
            for c in active_columns:
                cells.append(" ".join(w[4] for w in ws if (w[0]+w[2])/2>=c["x1"] and (w[0]+w[2])/2<c["x2"]).strip())
            if any(cells):out.append(cells)
        return out

    def export(self):
        if not self.doc:return
        if Workbook is None:
            messagebox.showerror("openpyxl Missing","Run INSTALL.bat included in this package.");return
        p=filedialog.asksaveasfilename(defaultextension=".xlsx",filetypes=[("Excel workbook","*.xlsx")])
        if not p:return
        try:
            wb=Workbook();ws=wb.active;ws.title="Data"
            heads=[c["name"] for c in self.columns];ws.append(heads);count=0
            for pi,page in enumerate(self.doc):
                for row in self.rows(page, page_index=pi):
                    vals=[]
                    for c,v in zip(heads,row):
                        if any(k in c.lower() for k in ["debit","credit","balance","amount","turnover"]):
                            vals.append(num(v) if v else None)
                        else:vals.append(v)
                    ws.append(vals);count+=1
            ws.freeze_panes="A2";ws.auto_filter.ref=ws.dimensions
            for i in range(1,len(heads)+1):
                ws.column_dimensions[get_column_letter(i)].width=min(40,max(10,len(heads[i-1])+2))
            wb.save(p);self.status.set(f"Excel workbook created: {count} rows")
            messagebox.showinfo("Complete",f"Excel workbook created with {count} rows.")
        except Exception as e:messagebox.showerror("Error",str(e))

if __name__=="__main__":
    root=tk.Tk();App(root);root.mainloop()
